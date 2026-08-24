import pandas as pd
from openpyxl import Workbook, load_workbook
from pathlib import Path

# TODO: Separate pandas and openpyxl operations, so that openpyxl handles workbook operations.
# And pandas task operations.


class ExcelStorage:
    COLUMNS = [
        "ID",
        "Project",
        "Task",
        "Status",
        "Priority",
        "Created",
        "Started",
        "Completed",
        "Next Action",
        "Notes",
    ]
    COLUMN_TYPES = {
        "ID": "Int64",
        "Project": "string",
        "Task": "string",
        "Status": "string",
        "Priority": "string",
        "Created": "string",
        "Started": "string",
        "Completed": "string",
        "Next Action": "string",
        "Notes": "string",
    }

    # TODO: Add handling so that overwrites do not destroy, now it is apparently ok. could be by project, so sheet per project.
    def __init__(self, filename="taskhandler.xlsx"):
        base_dir = Path(__file__).resolve().parent
        data_dir = base_dir / "data"
        data_dir.mkdir(parents=True, exist_ok=True)

        self.workbook = data_dir / filename

        self.create_workbook()

    def create_workbook(self):
        if self.workbook.exists():
            return

        wb = Workbook()

        sheet = wb.active
        sheet.title = "Inbox"  # or "Welcome"

        sheet.append(self.COLUMNS)

        wb.save(self.workbook)

    # OPENPYXL
    def get_projects(self):
        wb = load_workbook(self.workbook)
        return wb.sheetnames

    def create_project(self, name):
        wb = load_workbook(self.workbook)
        if name in wb.sheetnames:
            raise ValueError(f"Project '{name}' already exists.")

        sheet = wb.create_sheet(title=name)

        sheet.append(self.COLUMNS)

        wb.save(self.workbook)
    def rename_project(self, project, new_name):
        wb = load_workbook(self.workbook)
        wb_sheet = wb[project]
        if wb_sheet:
            wb_sheet.title = new_name
            wb.save(self.workbook)
        pass
    # PANDAS
    def load_tasks(self, project):
        tasks = pd.read_excel(
            self.workbook,
            sheet_name=project,
        )
        for column, dtype in self.COLUMN_TYPES.items():
            if column in tasks.columns:
                tasks[column] = tasks[column].astype(dtype)
        return tasks

    def save_tasks(self, tasks: pd.DataFrame, project):
        with pd.ExcelWriter(
            self.workbook, engine="openpyxl", mode="a", if_sheet_exists="replace"
        ) as writer:
            tasks.to_excel(writer, sheet_name=project, index=False)
